"""
1. Rebuild slides 5 & 6 with accurate career timeline (oldest to newest)
2. Add speaker notes to slides 5 & 6 with enhanced interview speech
"""

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
import lxml.etree as etree

def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

C_GOLD  = rgb('D4A843')
C_CYAN  = rgb('00D4E8')
C_BLUE  = rgb('3B8BEB')
C_GREEN = rgb('2DD48A')
C_WHITE = rgb('FFFFFF')
C_MUTED = rgb('8899BB')

SECTION_LABEL = (914400,  164592, 7315200, 320040)
COMPANY_BOX   = (914400,  411480, 8229600, 594360)
PAGE_NUM_BOX  = (10515600,347472, 1463040, 365760)

ROLE1_TOP  = 1417320
DATE_LEFT  = 457200;  DATE_W  = 1645920; DATE_H  = 320040
TITLE_LEFT = 2286000; TITLE_W = 9144000; TITLE_H = 347472
BULLET_L   = 2514600; BULLET_W= 9144000; BULLET_H= 274320
TITLE_GAP  = 50000
ROLE_GAP   = 380000


def purge_text(slide):
    NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    sp_tree = slide.shapes._spTree
    to_rm = []
    for child in list(sp_tree):
        if 'sp' not in child.tag:
            continue
        all_t = child.findall('.//{%s}t' % NS_A)
        if ''.join(t.text or '' for t in all_t).strip():
            to_rm.append(child)
    for el in to_rm:
        sp_tree.remove(el)


def add_tb(slide, left, top, w, h, text, size_pt, bold, color, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left=left, top=top, width=w, height=h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.color.rgb = color


def build(slide, page_label, company_title, roles):
    purge_text(slide)
    add_tb(slide, *SECTION_LABEL, 'PROFESSIONAL EXPERIENCE', 9, True, C_GOLD)
    add_tb(slide, *COMPANY_BOX,   company_title, 22, True, C_WHITE)
    add_tb(slide, *PAGE_NUM_BOX,  page_label, 10, False, C_MUTED, align=PP_ALIGN.RIGHT)

    y = ROLE1_TOP
    for role in roles:
        add_tb(slide, DATE_LEFT,  y, DATE_W,  DATE_H,  role['date'],  10, True, role['color'])
        add_tb(slide, TITLE_LEFT, y, TITLE_W, TITLE_H, role['title'], role.get('ts', 13), True, C_WHITE)
        y += TITLE_H + TITLE_GAP
        for b in role['bullets']:
            add_tb(slide, BULLET_L, y, BULLET_W, BULLET_H, b, 10.5, False, C_MUTED)
            y += BULLET_H
        y += ROLE_GAP


def set_notes(slide, notes_text):
    """Set speaker notes for a slide, creating notes slide if needed."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    # Clear existing
    for para in tf.paragraphs:
        for run in para.runs:
            run.text = ''
    # Set new text
    tf.text = notes_text


# ── Load presentation ───────────────────────────────────────────────────────
prs = Presentation('Eslam_Abdelazeem_Presentation_v5.pptx')

# ═══════════════════════════════════════════════════════════════════════════
#  SLIDE 5  —  Alexandria Electricity Distribution Company (2008-2020)
# ═══════════════════════════════════════════════════════════════════════════
build(
    prs.slides[4],
    '05 / 12',
    'Alexandria Electricity Distribution Company  |  2008 - 2020',
    [
        {
            "date":  "2008 - 2011",
            "color": C_GREEN,
            "title": "Operations & Maintenance Engineer  -  Medium Voltage Equipment",
            "ts":    12,
            "bullets": [
                "Responsible for operation and maintenance management of medium-voltage electrical equipment.",
                "Performed switching operations, inspections, and routine maintenance for distribution networks.",
            ]
        },
        {
            "date":  "2011 - 2015",
            "color": C_BLUE,
            "title": "Fault Restoration Management Engineer  -  Trouble Call System (TCS)",
            "ts":    12,
            "bullets": [
                "Managed fault restoration operations and supervised the Trouble Call System (TCS) control room.",
                "Coordinated emergency switching, outage management, and rapid network restoration activities.",
            ]
        },
        {
            "date":  "2015 - 2020",
            "color": C_GOLD,
            "title": "GIS Engineer  -  Asset Data Collection & Network Mapping",
            "ts":    12,
            "bullets": [
                "Collected and mapped coordinates of all electrical network assets (medium & low voltage) into ArcGIS.",
                "In 2018, joined the team evaluating technical specs for Lot 1 control centers (5 CC + 2 SCC) - GIS discipline.",
                "Participated in designing the electrical data model and network connectivity for the distribution control centers.",
            ]
        },
    ]
)

NOTES_SLIDE5 = """\
SPEAKER NOTES - Slide 5: Early & Mid Career (2008-2020)

"Good [morning / afternoon], and thank you sincerely for this opportunity. It is truly an honour to be here.

My name is Engineer Eslam Abdelazeem. I hold a Bachelor's degree in Computer & Systems Engineering from the Faculty of Engineering at Al-Azhar University.

Following my graduation and completion of military service in 2008, I joined Alexandria Electricity Distribution Company, where I began my career in Operations and Maintenance — responsible for managing medium-voltage electrical equipment, switching operations, and network reliability.

In 2011, I transitioned into Fault Restoration Management, operating within the Trouble Call System (TCS) control room — coordinating emergency response, rapid network restoration, and outage management across the distribution network.

In 2015, I moved into the Geographic Information Systems (GIS) department, where I was responsible for collecting, mapping, and registering the coordinates of all electrical network assets — at both medium and low voltage levels — into the ArcGIS platform, building a complete and accurate digital asset inventory.

In 2018, I was selected as part of a specialised team to evaluate the technical specifications for distribution control centers under Lot 1 — a major national project covering five distribution control centers and two supervisory control centers. In this role, I contributed directly to designing the electrical data model and defining the network connectivity architecture."
"""

set_notes(prs.slides[4], NOTES_SLIDE5)
print('Slide 5: content rebuilt + speaker notes added.')

# ═══════════════════════════════════════════════════════════════════════════
#  SLIDE 6  —  EEHC (2020-Present)
# ═══════════════════════════════════════════════════════════════════════════
build(
    prs.slides[5],
    '06 / 12',
    'Egyptian Electricity Holding Company (EEHC)  |  2020 - Present',
    [
        {
            "date":  "2020 - 2024",
            "color": C_GOLD,
            "title": "Senior Power GIS Engineer  -  Control Center Implementation",
            "ts":    12,
            "bullets": [
                "Moved to EEHC to lead implementation of distribution control center projects (GIS discipline).",
                "Oversaw all project milestones: data collection, validation, QA/QC, and migration into data warehouse.",
                "Executed FAT (Factory Acceptance Testing) and managed full integration between GIS, ADMS, and SCADA systems.",
                "Coordinated with Schneider Electric, Siemens, and GE for system delivery and commissioning.",
            ]
        },
        {
            "date":  "2024 - Present",
            "color": C_CYAN,
            "title": "Digital Command Center Manager",
            "ts":    13,
            "bullets": [
                "Promoted to Digital Command Center Manager overseeing platforms and dashboards at national level.",
                "Monitors and follows up on all platforms collecting operational data from 9 distribution companies across Egypt.",
                "Manages KPI dashboards, performance indicators, analytical reports, and data integration workflows.",
                "Drives digital transformation and automation across the Egyptian electricity distribution sector.",
            ]
        },
    ]
)

NOTES_SLIDE6 = """\
SPEAKER NOTES - Slide 6: EEHC & Digital Command Center (2020-Present)

"In 2020, I moved to the Egyptian Electricity Holding Company (EEHC) to lead the full implementation of the control center projects. I oversaw every phase — from data collection, validation, and QA/QC, through to data warehouse migration, Factory Acceptance Testing (FAT), and the complete integration of GIS, ADMS, and SCADA systems. I coordinated directly with international vendors including Schneider Electric, Siemens, and GE throughout this process.

In 2024, I was promoted to Digital Command Center Manager — responsible for supervising all platforms and dashboards that aggregate and monitor real-time operational data from nine electricity distribution companies spanning the entire Arab Republic of Egypt.

Throughout this journey of over 16 years, I have grown from a field engineer into a national-level technical manager — always driven by a passion for innovation, data-driven decision making, and building systems that serve millions of Egyptian citizens.

I believe my unique combination of deep field experience, enterprise GIS expertise, and digital platform management makes me well-positioned to contribute significantly to this role.

Thank you."
"""

set_notes(prs.slides[5], NOTES_SLIDE6)
print('Slide 6: content rebuilt + speaker notes added.')

prs.save('Eslam_Abdelazeem_Presentation_v5.pptx')
print('Saved successfully. Speaker notes now visible in Presenter View (Alt+F5).')
